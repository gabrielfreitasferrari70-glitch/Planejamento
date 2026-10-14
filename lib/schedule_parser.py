"""Parse schedules without storing uploaded documents or inventing missing dates."""
from collections import Counter
from datetime import datetime,date,time
from io import BytesIO,StringIO
import csv,hashlib,re,zipfile
import fitz
import openpyxl

def normalize(text):
    return re.sub(r'\s+', ' ', text).strip()


def boundaries(page):
    segments = []
    for drawing in page.get_drawings():
        if not drawing['fill'] or max(drawing['fill']) > .05:
            continue
        for item in drawing['items']:
            if item[0] == 're':
                rect = item[1]
                if rect.width > 50 and rect.height < 4:
                    segments.append(((rect.y0 + rect.y1) / 2, rect.width))
            elif item[0] == 'l':
                start, end = item[1:]
                if abs(start.y - end.y) < .3 and abs(start.x - end.x) > 50:
                    segments.append((start.y, abs(start.x - end.x)))
    groups = []
    for y, length in sorted(segments):
        if not groups or y - groups[-1][-1][0] > 2.5:
            groups.append([])
        groups[-1].append((y, length))
    return [sum(y * length for y, length in group) / sum(length for _, length in group)
            for group in groups if sum(length for _, length in group) > 350]


def cell_text(words, rect):
    lines = []
    selected = sorted((word for word in words
                       if rect.contains(fitz.Point((word[0]+word[2])/2, (word[1]+word[3])/2))),
                      key=lambda word: ((word[1]+word[3])/2, word[0]))
    for word in selected:
        center = (word[1]+word[3])/2
        if not lines or abs(center - lines[-1][0]) > 2:
            lines.append((center, []))
        lines[-1][1].append(word)
    return '\n'.join(' '.join(word[4] for word in sorted(line, key=lambda word: word[0]))
                     for _, line in lines)


def cell_color(pixmap, rect):
    # Compound PDF paths may enclose unfilled areas. Sample rendered backgrounds
    # rather than treating each path's enclosing rectangle as a filled cell.
    samples = []
    for y in [rect.y0+5, (rect.y0+rect.y1)/2, rect.y1-5]:
        for offset in [4, 5, 6]:
            samples.append(pixmap.pixel(round((rect.x0+offset)*2), round(y*2))[:3])
    color = Counter(samples).most_common(1)[0][0]
    return '#' + ''.join(f'{channel:02X}' for channel in color)


def event_kind(title):
    if re.search(r'prova|avaliaç|exame|teste|seminários?\s+[12]º\s*módulo|2[ªa]\s*chamada', title, re.I):
        return 'exam'
    if re.search(r'feriado|recesso|dia do servidor público', title, re.I):
        return 'break'
    if re.fullmatch(r'(?:livre(?: para estudo)?|horário livre)', title, re.I):
        return 'free'
    if re.search(r'tutoria', title, re.I):
        return 'tutorial'
    if re.search(r'interação comunitária', title, re.I):
        return 'community'
    if re.search(r'acolhimento', title, re.I):
        return 'welcome'
    return 'class'


def subject(title):
    if ':' in title:
        return title.split(':', 1)[0].strip()
    for label in ['Anatomia', 'Antropologia', 'Interação Comunitária']:
        if re.search(re.escape(label), title, re.I):
            return label
    return 'Geral'


def date_value(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    text = str(value or '').strip()
    for fmt in ['%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y']:
        try: return datetime.strptime(text, fmt).date().isoformat()
        except ValueError: pass
    return None


def clock(value):
    if isinstance(value, (datetime, time)): return value.strftime('%H:%M')
    match = re.fullmatch(r'(\d{1,2}):(\d{2})(?::00)?', str(value or '').strip())
    if match and int(match[1]) < 24 and int(match[2]) < 60:
        return match[1].zfill(2) + ':' + match[2]
    raise ValueError('Há um horário inválido na planilha. Use HH:MM.')


def record(day, start, end, text, professor='', color=None, kind=None, discipline=None):
    if not day: raise ValueError('Há uma data inválida. Use DD/MM/AAAA ou AAAA-MM-DD.')
    start, end = clock(start), clock(end)
    if start >= end: raise ValueError('O horário final precisa ser depois do inicial.')
    text = str(text or '').strip()
    match = re.search(r'\b(?:Prof(?:essor(?:a)?|[saª])?\.?|Docente)\s+', text, re.I)
    if match and not professor:
        professor, text = normalize(text[match.start():]), text[:match.start()]
    title = normalize(text).rstrip(' -–—')
    if not title: raise ValueError('Há uma atividade sem título.')
    result = {'date':day,'from':start,'to':end,'title':title[:500],
              'subject':str(discipline or subject(title))[:100],
              'type':kind if kind in ['class','exam','break','free','tutorial','community','welcome'] else event_kind(title),
              'professor':normalize(str(professor or ''))[:200],'origin':'import'}
    if re.fullmatch(r'#[0-9a-fA-F]{6}', color or ''): result['sheetColor'] = color
    return result


def parse_rows(rows):
    # Both a table with named columns and a weekly grid with date headers.
    rows = list(rows)
    if len(rows) > 10000: raise ValueError('A planilha ultrapassa 10 mil linhas.')
    alias = {'data':'date','date':'date','inicio':'from','hora inicio':'from','from':'from',
             'fim':'to','hora fim':'to','to':'to','titulo':'title','atividade':'title','title':'title',
             'disciplina':'subject','subject':'subject','professor':'professor','tipo':'type','type':'type','cor':'color'}
    for index, row in enumerate(rows[:20]):
        names = [alias.get(normalize(str(v or '')).lower().translate(str.maketrans('íúãç','iuac'))) for v in row]
        if not {'date','from','to','title'}.issubset(set(names)): continue
        result=[]
        for values in rows[index+1:]:
            if not any(v is not None and str(v).strip() for v in values): continue
            obj={key:values[i] for i,key in enumerate(names) if key and i<len(values)}
            result.append(record(date_value(obj.get('date')),obj.get('from'),obj.get('to'),obj.get('title'),
                                 obj.get('professor'),obj.get('color'),obj.get('type'),obj.get('subject')))
        return result
    result=[];headers=None
    time_range=re.compile(r'(\d{1,2}:\d{2})\s*[-–—]\s*(\d{1,2}:\d{2})')
    for row in rows:
        days=[date_value(v) for v in row]
        if sum(bool(d) for d in days)>=2: headers=days;continue
        if not headers: continue
        timed=next(((i,time_range.fullmatch(str(v or '').strip())) for i,v in enumerate(row) if time_range.fullmatch(str(v or '').strip())),None)
        if not timed: continue
        index,match=timed
        for column,day in enumerate(headers):
            if day and column!=index and column<len(row) and str(row[column] or '').strip():
                result.append(record(day,match[1],match[2],row[column]))
    return result


def parse_pdf(body):
    result=[]
    with fitz.open(stream=body,filetype='pdf') as document:
        if document.needs_pass: raise ValueError('Remova a senha do PDF antes de importar.')
        if len(document)>40: raise ValueError('Use um PDF com até 40 páginas.')
        for page in document:
            words=page.get_text('words')
            dated=[w for w in words if date_value(w[4])]
            headers=[]
            for word in sorted(dated,key=lambda w:((w[1]+w[3])/2,w[0])):
                center=(word[1]+word[3])/2
                if not headers or abs(center-headers[-1][0])>3: headers.append((center,[]))
                headers[-1][1].append(word)
            headers=[(y,sorted(columns,key=lambda w:w[0])) for y,columns in headers if len(columns)>=2]
            timed=[(w,re.fullmatch(r'(\d{1,2}:\d{2})\s*[-–—]\s*(\d{1,2}:\d{2})',w[4])) for w in words]
            timed=[(w,m) for w,m in timed if m]
            if not headers or not timed: raise ValueError('Não reconheci uma grade com datas e horários neste PDF. PDFs escaneados precisam ser convertidos para Excel/CSV.')
            if page.rect.width*page.rect.height>2000000: raise ValueError('O tamanho das páginas do PDF é muito grande.')
            edges=boundaries(page)
            pixmap=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
            for word,match in timed:
                y=(word[1]+word[3])/2
                candidates=[header for header in headers if header[0]<y]
                if not candidates: continue
                _,columns=candidates[-1]
                centers=[(w[0]+w[2])/2 for w in columns]
                width=(centers[-1]-centers[0])/(len(centers)-1)
                above=[edge for edge in edges if edge<y];below=[edge for edge in edges if edge>y]
                if not above or not below: raise ValueError('Não consegui separar as células da grade do PDF. Importe a versão Excel/CSV.')
                top,bottom=max(above),min(below)
                for column,header in enumerate(columns):
                    rect=fitz.Rect(max(0,centers[column]-width/2+1),top+1,min(page.rect.width,centers[column]+width/2-1),bottom-1)
                    text=cell_text(words,rect)
                    if text: result.append(record(date_value(header[4]),match[1],match[2],text,color=cell_color(pixmap,rect)))
    return result


def analyze(body,filename):
    extension=filename.lower().rsplit('.',1)[-1]
    if extension=='pdf': rows=parse_pdf(body)
    elif extension=='csv':
        try: text=body.decode('utf-8-sig')
        except UnicodeDecodeError: text=body.decode('cp1252')
        try: dialect=csv.Sniffer().sniff(text[:10000],delimiters=',;\t')
        except csv.Error: dialect=csv.excel
        rows=parse_rows(csv.reader(StringIO(text),dialect))
    elif extension=='xlsx':
        with zipfile.ZipFile(BytesIO(body)) as archive:
            if sum(item.file_size for item in archive.infolist())>20000000: raise ValueError('A planilha descompactada é muito grande.')
        workbook=openpyxl.load_workbook(BytesIO(body),read_only=True,data_only=True)
        try:
            rows=[]
            for sheet in workbook:
                if sheet.max_column>100 or sheet.max_row>10000: raise ValueError('Use uma planilha com até 100 colunas e 10 mil linhas.')
                rows+=parse_rows(sheet.iter_rows(values_only=True))
        finally: workbook.close()
    else: raise ValueError('Use um arquivo PDF, XLSX ou CSV.')
    if not rows: raise ValueError('Nenhuma atividade foi encontrada. Confira se a planilha contém datas, horários e títulos.')
    unique={tuple(row.get(k) for k in ['date','from','to','title','professor','sheetColor']):row for row in rows}
    merged=[]
    for row in sorted(unique.values(),key=lambda x:(x['date'],x['from'])):
        previous=merged[-1] if merged else None
        if previous and previous['date']==row['date'] and previous['to']==row['from'] and all(previous.get(k)==row.get(k) for k in ['title','subject','type','professor','sheetColor']): previous['to']=row['to']
        else: merged.append(dict(row))
    if len(merged)>2000: raise ValueError('O cronograma ultrapassa 2 mil atividades.')
    for row in merged:
        identity='|'.join(row[k] for k in ['date','from','to','title'])
        row['id']='schedule-'+hashlib.sha256(identity.encode()).hexdigest()[:24]
    first,last=merged[0]['date'],merged[-1]['date']
    counts=dict(Counter(row['type'] for row in merged))
    names={'class':'Aulas','exam':'Avaliações','break':'Feriados','free':'Horário livre','tutorial':'Tutorias','community':'Interação Comunitária','welcome':'Acolhimento'}
    legend=[]
    for kind,label in names.items():
        colors=Counter(row['sheetColor'] for row in merged if row['type']==kind and row.get('sheetColor'))
        if colors: legend.append({'color':colors.most_common(1)[0][0],'name':label})
    metadata={'filename':filename[:150],'start':first,'end':last,'events':len(merged),'counts':counts,'legend':legend,
              'semester':first[:4]+'/'+('1' if int(first[5:7])<=6 else '2')}
    return {'events':merged,'source':metadata}


